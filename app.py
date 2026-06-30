import re
import html
import streamlit as st
import pandas as pd
from google import genai
from google.genai import types
from datetime import datetime
from PyPDF2 import PdfReader
from docx import Document
from PIL import Image
from fpdf import FPDF
from fpdf.enums import XPos, YPos
from supabase import create_client, Client

# Konstanta newline char(10) - dipakai agar tidak ada escape sequence di dalam kode
NL = chr(10)

# ══════════════════════════════════════════════════════
# KONFIGURASI HALAMAN
# ══════════════════════════════════════════════════════
st.set_page_config(
    page_title="Buat CV Mu — AI Career Optimizer",
    page_icon="📄",
    layout="wide",
    initial_sidebar_state="collapsed"
)

# ══════════════════════════════════════════════════════
# LOAD CONFIG
# ══════════════════════════════════════════════════════
def load_config():
    try:
        gemini_key = st.secrets.get("GEMINI_API_KEY", None)
        supabase_url = st.secrets.get("SUPABASE_URL", None)
        supabase_key = st.secrets.get("SUPABASE_KEY", None)
    except Exception:
        gemini_key = supabase_url = supabase_key = None
    return gemini_key, supabase_url, supabase_key

gemini_key, supabase_url, supabase_key = load_config()

@st.cache_resource
def init_supabase(url, key):
    if url and key:
        return create_client(url, key)
    return None

supabase: Client = init_supabase(supabase_url, supabase_key)

# ══════════════════════════════════════════════════════
# EKSTRAKSI DOKUMEN
# ══════════════════════════════════════════════════════
def extract_pdf_text(file) -> str:
    try:
        reader = PdfReader(file)
        return NL.join(page.extract_text() or "" for page in reader.pages).strip()
    except Exception as e:
        return f"[ERROR PDF: {e}]"

def extract_docx_text(file) -> str:
    try:
        doc = Document(file)
        return NL.join(p.text for p in doc.paragraphs if p.text.strip()).strip()
    except Exception as e:
        return f"[ERROR DOCX: {e}]"

def extract_cv_text(uploaded_file) -> str:
    name = uploaded_file.name.lower()
    if name.endswith(".pdf"):
        return extract_pdf_text(uploaded_file)
    elif name.endswith(".docx"):
        return extract_docx_text(uploaded_file)
    return "[Format tidak didukung]"

# ══════════════════════════════════════════════════════
# EXPORT PDF
# ══════════════════════════════════════════════════════
def _strip_md(text: str) -> str:
    text = text.replace("**", "").replace("*", "")     # buang penanda bold/italic
    text = re.sub("(?m)^#{1,6} *", "", text)           # buang penanda heading
    text = text.replace("•", "-")
    return text

def _latinize(text: str) -> str:
    repl = {"—": "-", "–": "-", "’": "'", "‘": "'", "“": '"', "”": '"', "…": "...", "•": "-"}
    for k, v in repl.items():
        text = text.replace(k, v)
    return text.encode("latin-1", "ignore").decode("latin-1")

def make_pdf(title: str, body: str) -> bytes:
    pdf = FPDF(format="A4")
    pdf.add_page()
    pdf.set_auto_page_break(auto=True, margin=15)
    pdf.set_font("Helvetica", "B", 15)
    pdf.multi_cell(0, 9, _latinize(title), new_x=XPos.LMARGIN, new_y=YPos.NEXT)
    pdf.ln(2)
    pdf.set_font("Helvetica", size=11)
    for line in _strip_md(body).split(NL):
        line = _latinize(line.rstrip())
        if line.strip() == "":
            pdf.ln(3)
        else:
            pdf.multi_cell(0, 6, line, new_x=XPos.LMARGIN, new_y=YPos.NEXT)
    return bytes(pdf.output())

BULLET = chr(183)  # middot, aman untuk font latin-1

def _is_full_bold(s):
    return s.startswith("**") and s.endswith("**") and len(s) > 4 and "|" not in s

def _is_bullet(s):
    return s.startswith("- ") or s.startswith(chr(8226)) or (s.startswith("* ") and not s.startswith("**"))

def make_cv_pdf(text: str) -> bytes:
    pdf = FPDF(format="A4")
    pdf.set_margins(16, 14, 16)
    pdf.set_auto_page_break(auto=True, margin=14)
    pdf.add_page()
    epw = pdf.epw
    expect_contact = False

    for raw in text.split(NL):
        s = raw.strip()
        if s == "":
            if not expect_contact:
                pdf.ln(2)
            continue

        # NAMA (judul besar di tengah)
        if s.startswith("# "):
            pdf.set_font("Helvetica", "B", 24); pdf.set_text_color(20, 20, 20)
            pdf.multi_cell(0, 11, _latinize(s[2:].replace("**", "").strip()).upper(), align="C", new_x=XPos.LMARGIN, new_y=YPos.NEXT)
            expect_contact = True
            continue

        # KONTAK (baris tepat setelah nama)
        if expect_contact:
            expect_contact = False
            pdf.set_font("Helvetica", "", 10); pdf.set_text_color(70, 70, 70)
            for part in s.replace("**", "").split("|"):
                if part.strip():
                    pdf.multi_cell(0, 5, _latinize(part.strip()), align="C", new_x=XPos.LMARGIN, new_y=YPos.NEXT)
            pdf.ln(1)
            continue

        # BARIS JUDUL + TANGGAL (kiri/kanan, pakai pipe)
        if "|" in s:
            left, right = s.split("|", 1)
            pdf.set_text_color(20, 20, 20); pdf.set_font("Helvetica", "B", 10.5)
            pdf.cell(epw * 0.70, 5.5, _latinize(left.replace("**", "").replace("#", "").strip()), new_x=XPos.RIGHT, new_y=YPos.TOP)
            pdf.cell(epw * 0.30, 5.5, _latinize(right.replace("**", "").strip()), align="R", new_x=XPos.LMARGIN, new_y=YPos.NEXT)
            continue

        # HEADER SECTION (full bold, di tengah + garis bawah)
        if _is_full_bold(s):
            pdf.ln(3); pdf.set_font("Helvetica", "B", 12); pdf.set_text_color(20, 20, 20)
            pdf.multi_cell(0, 6.5, _latinize(s.replace("**", "").strip()).upper(), align="C", new_x=XPos.LMARGIN, new_y=YPos.NEXT)
            y = pdf.get_y() + 0.5
            pdf.set_draw_color(60, 60, 60); pdf.set_line_width(0.4)
            pdf.line(pdf.l_margin, y, pdf.l_margin + epw, y)
            pdf.ln(2.5)
            continue

        # BULLET
        if _is_bullet(s):
            txt = s.lstrip("-* " + chr(8226)).strip()
            pdf.set_font("Helvetica", "", 10); pdf.set_text_color(45, 45, 45)
            pdf.set_x(pdf.l_margin + 2)
            pdf.cell(4, 5, BULLET, new_x=XPos.RIGHT, new_y=YPos.TOP)
            pdf.multi_cell(epw - 6, 5, _latinize(txt), new_x=XPos.LMARGIN, new_y=YPos.NEXT)
            continue

        # PARAGRAF biasa
        pdf.set_font("Helvetica", "", 10); pdf.set_text_color(45, 45, 45)
        pdf.multi_cell(0, 5, _latinize(s.replace("**", "")), new_x=XPos.LMARGIN, new_y=YPos.NEXT)

    return bytes(pdf.output())

# ══════════════════════════════════════════════════════
# SUPABASE
# ══════════════════════════════════════════════════════
def save_to_supabase(company, position, status, full_output, cover_snippet):
    if not supabase:
        return False
    try:
        supabase.table("cv_applications").insert({
            "timestamp": datetime.now().isoformat(),
            "company_name": company,
            "target_position": position,
            "status": status,
            "cover_letter_snippet": cover_snippet[:200],
            "full_output": full_output[:5000]
        }).execute()
        return True
    except Exception as e:
        st.warning(f"Log gagal disimpan: {e}")
        return False

def fetch_history() -> pd.DataFrame:
    if not supabase:
        return pd.DataFrame()
    try:
        res = (
            supabase.table("cv_applications")
            .select("timestamp,company_name,target_position,status,cover_letter_snippet")
            .order("timestamp", desc=True)
            .limit(50)
            .execute()
        )
        return pd.DataFrame(res.data)
    except Exception:
        return pd.DataFrame()

# ══════════════════════════════════════════════════════
# SYSTEM PROMPT
# ══════════════════════════════════════════════════════
SYSTEM_PROMPT = """
Kamu adalah mesin backend dari aplikasi karir bernama "Buat CV Mu".

Kamu berperan sebagai tiga entitas sekaligus:
1. HR Psikolog Senior dengan 20+ tahun pengalaman membaca manusia dari CV
2. ATS Specialist yang tahu persis cara mengalahkan sistem filtering perusahaan Fortune 500
3. Personal Branding Strategist yang tahu cara membuat rekruter berhenti scroll

Tugasmu adalah memproses bahan lamaran kerja pengguna dengan logika bisnis yang KETAT.

════════════════════════════════════════════
STEP 1: GATEKEEPER — ANALISIS KECOCOKAN
════════════════════════════════════════════

Evaluasi kompatibilitas CV dengan posisi yang dituju secara KRITIS. Jangan jadi yes-man.

- CASE A (SANGAT TIDAK RELEVAN): Jika kualifikasi benar-benar tidak bisa dijembatani (misal: lulusan desain grafis melamar dokter spesialis), TOLAK. Output HANYA:
  ! [DITOLAK OLEH AI GATEKEEPER]
  [Penjelasan tajam dalam Bahasa Indonesia: mengapa ditolak, posisi apa yang lebih sesuai]
  Tidak ada output lain.

- CASE B (ADA GAP TAPI BISA DIJEMBATANI): Lanjut ke Step 2. Gali transferable skills, frame ulang pengalaman secara strategis.

- CASE C (COCOK): Lanjut ke Step 2. Maksimalkan semua potensi.

════════════════════════════════════════════
STEP 2: OUTPUT (Hanya Case B atau C)
════════════════════════════════════════════

Gunakan PERSIS header-header berikut. Jangan ubah format header.

---SECTION:COMMENTARY---

Tulis dalam BAHASA INDONESIA. Kamu adalah HR Psikolog Senior yang membedah profil ini secara mendalam.

Struktur analisis:
**🔍 Pembacaan Karakter & Pola Karir:**
[Baca pola karir, motivasi tersembunyi, karakter kepribadian yang terbaca dari CV. Bukan sekadar ringkasan — tapi insight psikologis. Apakah ini orang yang driven by money, status, impact, atau stabilitas? Apa yang tidak ditulis tapi terbaca?]

**⚡ Kekuatan Tersembunyi yang Perlu Ditonjolkan:**
[Temukan 3-5 kekuatan non-obvious yang sering diabaikan kandidat tapi sangat menarik bagi rekruter untuk posisi ini]

**🚧 Gap & Red Flags yang Harus Diatasi:**
[Jujur dan tajam. Apa yang akan langsung dipertanyakan rekruter? Bagaimana cara menjawabnya sebelum ditanya?]

**🎯 VERDICT — Rekomendasi Strategis:**
[Berikan verdict tegas: SANGAT DIREKOMENDASIKAN / DIREKOMENDASIKAN DENGAN CATATAN / PERTIMBANGKAN ULANG]
[Sertakan alasan konkret dan langkah taktis yang harus dilakukan SEBELUM melamar jika ada gap]

---SECTION:ATS_SCORE---

Hitung skor kecocokan CV terhadap posisi & deskripsi pekerjaan secara realistis (0-100). Gunakan format PERSIS ini (jangan tambah penjelasan lain):
SCORE: [angka 0-100 tanpa simbol]
RINGKASAN: [satu kalimat singkat alasan skor tersebut]
KEYWORD_COCOK: [keyword dari CV yang sudah relevan dengan posisi, dipisah koma, maksimal 10]
KEYWORD_HILANG: [keyword penting dari lowongan yang BELUM ada di CV, dipisah koma, maksimal 10]

---SECTION:CV_ENGLISH---

Buat ATS-Optimized CV dalam Bahasa Inggris. WAJIB ikuti format baris berikut PERSIS (supaya bisa dirender rapi menjadi PDF):

# [Nama Lengkap Kandidat]
[Nomor Telepon] | [Email] | [Kota/Lokasi]

**PROFESSIONAL SUMMARY**
[3-4 kalimat memukau dalam bentuk paragraf, spesifik ke posisi, pakai angka/pencapaian/value proposition. Bukan bullet.]

**CORE COMPETENCIES**
- [skill/keyword ATS 1]
- [skill/keyword ATS 2]
(total 8-15 item, masing-masing satu baris diawali tanda "- ")

**PROFESSIONAL EXPERIENCE**
**[Job Title] - [Nama Perusahaan]** | [Bln Thn - Bln Thn]
- [action verb + hal yang dikerjakan + hasil/metrik]
- [pencapaian lain]
(ulangi blok ini untuk tiap posisi. Baris jabatan WAJIB format: **Judul - Perusahaan** | Tanggal)

**EDUCATION & CERTIFICATIONS**
**[Nama Institusi/Penerbit]** | [Tahun]
[Gelar/Program/Nama Sertifikat]
(ulangi untuk tiap entri)

**KEY ACHIEVEMENTS**
- [pencapaian terbaik dengan angka konkret]
(3-5 item)

---SECTION:CV_INDONESIA---

Versi Bahasa Indonesia (adaptasi natural, bukan terjemahan kaku; tetap ATS-friendly untuk JobStreet/Glints/LinkedIn ID). WAJIB ikuti format baris yang sama:

# [Nama Lengkap Kandidat]
[Nomor Telepon] | [Email] | [Kota/Lokasi]

**RINGKASAN PROFESIONAL**
[paragraf 3-4 kalimat, diksi profesional Indonesia yang modern]

**KOMPETENSI INTI**
- [skill/keyword 1]
- [skill/keyword 2]
(8-15 item, masing-masing satu baris diawali "- ")

**PENGALAMAN KERJA**
**[Jabatan] - [Nama Perusahaan]** | [Bln Thn - Bln Thn]
- [kata kerja aksi + hasil + metrik]
- [pencapaian lain]
(ulangi tiap posisi. Baris jabatan WAJIB format: **Jabatan - Perusahaan** | Tanggal)

**PENDIDIKAN & SERTIFIKASI**
**[Nama Institusi]** | [Tahun]
[Gelar/Program/Sertifikat]
(ulangi tiap entri)

**PENCAPAIAN UTAMA**
- [pencapaian dengan angka konkret]
(3-5 item)

---SECTION:COVER_LETTER---

Tulis cover letter dalam BAHASA INDONESIA.
Tone: Profesional tapi tidak kaku. Modern, percaya diri, sedikit personal — seperti orang yang tahu nilainya sendiri.
WAJIB di bawah 500 karakter total (termasuk spasi).
Mulai langsung dengan hook — JANGAN mulai dengan "Dengan hormat" atau "Saya yang bertanda tangan".
Hitung karakter dengan teliti sebelum output.

---SECTION:INTERVIEW---

**Pertanyaan 1 — "Ceritakan tentang diri Anda"**
*(Frame khusus untuk perusahaan dan posisi ini)*
Rangkaian Jawaban:
• [Poin strategis 1 — opening hook yang memorable]
• [Poin strategis 2 — bukti track record relevan]
• [Poin strategis 3 — koneksi ke value perusahaan ini]
• [Poin strategis 4 — closing statement yang confident]

**Pertanyaan 2 — [Pertanyaan behavioral relevan dengan posisi]**
Rangkaian Jawaban:
• [Framework STAR: Situation]
• [Framework STAR: Task & Action]
• [Framework STAR: Result dengan angka]

**Pertanyaan 3 — [Pertanyaan tentang kelemahan/tantangan]**
Rangkaian Jawaban:
• [Cara menjawab yang jujur tapi strategis]
• [Bukti self-awareness]
• [Growth mindset yang ditunjukkan]

**Pertanyaan 4 — [Pertanyaan teknikal/situasional spesifik posisi]**
Rangkaian Jawaban:
• [Jawaban berbasis pengalaman nyata]
• [Pendekatan problem-solving]
• [Differentiator dari kandidat lain]

**Pertanyaan 5 — "Mengapa Anda ingin bergabung dengan perusahaan kami?"**
Rangkaian Jawaban:
• [Riset spesifik tentang perusahaan yang harus disebutkan]
• [Koneksi antara company goals dan personal goals]
• [Apa yang bisa kamu bawa yang orang lain tidak bisa]

---SECTION:DBLOG---

| Timestamp | Perusahaan | Posisi | Status | Cuplikan Cover Letter |
|-----------|-----------|--------|--------|----------------------|
| {DATETIME} | {COMPANY} | {POSITION} | APPROVED | [100 karakter pertama cover letter] |

ATURAN KRITIS:
- Jangan pernah keluar dari format section di atas
- Cover letter HARUS di bawah 500 karakter. Hitung ulang sebelum output.
- Semua section harus terisi penuh — jangan ada yang dipersingkat
- Analisis Commentary harus MENDALAM, bukan permukaan
- CV harus benar-benar outstanding — bukan template generik
"""

# ══════════════════════════════════════════════════════
# PARSER
# ══════════════════════════════════════════════════════
def parse_sections(raw: str) -> dict:
    sections = {
        "commentary": "", "ats_score": "", "cv_english": "", "cv_indonesia": "",
        "cover_letter": "", "interview": "", "dblog": "", "rejected": False, "raw": raw
    }

    if "[DITOLAK OLEH AI GATEKEEPER]" in raw:
        sections["rejected"] = True
        sections["commentary"] = raw
        return sections

    markers = [
        ("commentary",   "---SECTION:COMMENTARY---"),
        ("ats_score",    "---SECTION:ATS_SCORE---"),
        ("cv_english",   "---SECTION:CV_ENGLISH---"),
        ("cv_indonesia", "---SECTION:CV_INDONESIA---"),
        ("cover_letter", "---SECTION:COVER_LETTER---"),
        ("interview",    "---SECTION:INTERVIEW---"),
        ("dblog",        "---SECTION:DBLOG---"),
    ]

    for i, (key, marker) in enumerate(markers):
        start = raw.find(marker)
        if start == -1:
            continue
        start += len(marker)
        if i + 1 < len(markers):
            end = raw.find(markers[i + 1][1])
            sections[key] = raw[start:end].strip() if end != -1 else raw[start:].strip()
        else:
            sections[key] = raw[start:].strip()

    return sections

def parse_ats(text: str):
    score = 0
    m = re.search("SCORE: *([0-9]{1,3})", text)
    if m:
        score = max(0, min(100, int(m.group(1))))

    def grab(label):
        mm = re.search(label + ": *(.+)", text)
        if not mm:
            return []
        return [x.strip() for x in re.split("[," + NL + "]", mm.group(1)) if x.strip()][:12]

    matched = grab("KEYWORD_COCOK")
    missing = grab("KEYWORD_HILANG")
    sm = re.search("RINGKASAN: *(.+)", text)
    summary = sm.group(1).strip() if sm else ""
    return score, summary, matched, missing

# ══════════════════════════════════════════════════════
# CORE ANALISIS
# ══════════════════════════════════════════════════════
def analyze_application(cv_text, company_name, position, extra_info, job_image=None, api_key=None) -> str:
    key = api_key or gemini_key
    if not key:
        return "[ERROR] Gemini API Key tidak ditemukan."

    client = genai.Client(api_key=key)

    user_prompt = f"""
DATETIME SEKARANG: {datetime.now().strftime("%Y-%m-%d %H:%M:%S")}
PERUSAHAAN TARGET: {company_name}
POSISI YANG DILAMAR: {position}
INFO TAMBAHAN / DESKRIPSI PEKERJAAN: {extra_info if extra_info else "Tidak ada info tambahan."}

TEKS CV KANDIDAT:
{cv_text}
"""

    contents = []
    if job_image is not None:
        job_image.seek(0)
        img = Image.open(job_image)
        contents.append(img)
        contents.append("Analisis gambar screenshot lowongan kerja ini. Ekstrak semua kualifikasi, requirement, keyword, dan ekspektasi perusahaan dari gambar tersebut.")

    contents.append(user_prompt)

    try:
        response = client.models.generate_content(
            model="gemini-2.5-flash",
            contents=contents,
            config=types.GenerateContentConfig(
                system_instruction=SYSTEM_PROMPT,
                temperature=0.5,
                max_output_tokens=32768,
                thinking_config=types.ThinkingConfig(thinking_budget=8192),
            ),
        )
        return response.text or "[ERROR] Model tidak mengembalikan teks (kemungkinan diblokir filter keamanan)."
    except Exception as e:
        return f"[ERROR dari Gemini API]: {str(e)}"

# ══════════════════════════════════════════════════════
# SIDEBAR + TEMA (harus sebelum inject CSS)
# ══════════════════════════════════════════════════════
with st.sidebar:
    st.markdown("### ⚙️ Konfigurasi")
    dark_mode = st.toggle("🌙 Mode Gelap", value=False, key="dark_mode")
    st.markdown("---")
    manual_api_key = st.text_input("Gemini API Key (Fallback)", type="password")
    st.markdown("---")
    st.caption("v5.2 · CV PDF Rapi · google-genai · Neumorphic")
    st.success("✓ Supabase terhubung") if supabase else st.warning("⚠️ Supabase tidak terhubung")

theme = "dark" if dark_mode else "light"

# ══════════════════════════════════════════════════════
# PALET TEMA → CSS VARIABLES
# ══════════════════════════════════════════════════════
PALETTES = {
    "light": {"bg": "#e6e7f1", "text": "#3a3a55", "soft": "#54547a", "muted": "#8a8aa8",
              "sd": "#c4c5d4", "sl": "#ffffff", "sd2": "#d2d3e0",
              "accent": "#7c3aed", "accent2": "#6024c9", "phold": "#9a9ab5"},
    "dark":  {"bg": "#20203a", "text": "#e8e8f7", "soft": "#bcbce0", "muted": "#8585aa",
              "sd": "#15152a", "sl": "#2c2c52", "sd2": "#191932",
              "accent": "#a78bfa", "accent2": "#7c3aed", "phold": "#7070a0"},
}
p = PALETTES[theme]
st.markdown(f"""<style>:root {{
  --bg:{p['bg']}; --text:{p['text']}; --soft:{p['soft']}; --muted:{p['muted']};
  --sd:{p['sd']}; --sl:{p['sl']}; --sd2:{p['sd2']};
  --accent:{p['accent']}; --accent2:{p['accent2']}; --phold:{p['phold']};
}}</style>""", unsafe_allow_html=True)

# ══════════════════════════════════════════════════════
# CSS UTAMA (neumorphic, pakai variabel tema)
# ══════════════════════════════════════════════════════
st.markdown("""
<style>
    @import url('https://fonts.googleapis.com/css2?family=Poppins:wght@400;500;600;700;800&display=swap');

    .stApp { background: var(--bg); font-family: 'Poppins', sans-serif; }
    html, body, [class*="css"], p, span, label, div { font-family: 'Poppins', sans-serif; color: var(--text); }
    [data-testid="stSidebar"] { background: var(--bg); }
    #MainMenu, footer, header { visibility: hidden; }
    .block-container { padding-top: 2rem; max-width: 1200px; }

    /* ===== HERO PANEL ===== */
    .neo-hero {
        background: linear-gradient(150deg, #1b1b3a 0%, #322a6e 55%, #1f1f44 100%);
        border-radius: 28px; padding: 1.8rem 1.6rem; margin-bottom: 1.4rem;
        border: 1px solid rgba(167,139,250,0.25);
        box-shadow: 10px 10px 24px var(--sd), -10px -10px 24px var(--sl);
        animation: fadeUp 0.5s ease both;
    }
    .neo-hero h1 { color: #fff !important; font-size: 1.55rem; font-weight: 800; margin: 0; letter-spacing: -0.5px; }
    .neo-hero p { color: #c3b8ff !important; margin: 0.35rem 0 0; font-size: 0.82rem; }
    .neo-hero .pill {
        display:inline-block; margin-top: 0.9rem; padding: 0.3rem 0.9rem;
        background: rgba(255,255,255,0.10); border-radius: 30px;
        color:#e3ddff !important; font-size:0.72rem; font-weight:500;
    }

    .neo-label { font-weight: 700; font-size: 1rem; color: var(--text) !important; margin: 0.6rem 0 0.8rem; }

    /* ===== INPUTS ===== */
    .stTextInput input, .stTextArea textarea {
        background: var(--bg) !important; border: none !important; border-radius: 14px !important;
        box-shadow: inset 4px 4px 9px var(--sd), inset -4px -4px 9px var(--sl) !important;
        color: var(--text) !important; padding: 0.65rem 0.9rem !important;
    }
    .stTextInput input::placeholder, .stTextArea textarea::placeholder { color: var(--phold) !important; }
    .stTextInput label, .stTextArea label, .stFileUploader label {
        font-weight: 600 !important; font-size: 0.82rem !important; color: var(--soft) !important;
    }

    [data-testid="stFileUploader"] section {
        background: var(--bg) !important; border: 2px dashed var(--sd) !important; border-radius: 16px !important;
        box-shadow: inset 3px 3px 7px var(--sd), inset -3px -3px 7px var(--sl);
    }
    [data-testid="stFileUploader"] section * { color: var(--soft) !important; }

    /* ===== BUTTONS ===== */
    .stButton > button {
        background: var(--bg); color: var(--soft); border: none; border-radius: 16px;
        padding: 0.7rem 1rem; font-weight: 700; font-size: 0.9rem; width: 100%;
        box-shadow: 6px 6px 13px var(--sd), -6px -6px 13px var(--sl); transition: all 0.18s ease;
    }
    .stButton > button:hover { color: var(--accent); transform: translateY(-2px); box-shadow: 8px 8px 16px var(--sd), -8px -8px 16px var(--sl); }
    .stButton > button:active { box-shadow: inset 4px 4px 9px var(--sd), inset -4px -4px 9px var(--sl); }
    .stButton > button[kind="primary"] {
        background: linear-gradient(145deg, var(--accent), var(--accent2)); color: #fff !important;
        box-shadow: 6px 6px 14px var(--sd), -6px -6px 14px var(--sl), inset 1px 1px 2px rgba(255,255,255,0.25);
    }
    .stButton > button[kind="primary"]:hover { color:#fff !important; transform: translateY(-2px); }

    .stDownloadButton > button {
        background: var(--bg); color: var(--accent); border:none; border-radius:14px;
        font-weight:600; font-size:0.8rem; box-shadow: 4px 4px 9px var(--sd), -4px -4px 9px var(--sl);
    }
    .stDownloadButton > button:hover { transform: translateY(-1px); }

    /* ===== CONTENT CARD ===== */
    .neo-card {
        background: var(--bg); border-radius: 24px; padding: 1.6rem 1.8rem; margin-top: 0.4rem;
        box-shadow: inset 5px 5px 12px var(--sd), inset -5px -5px 12px var(--sl);
    }
    [data-testid="stVerticalBlockBorderWrapper"] {
        background: var(--bg); border: none !important; border-radius: 24px; padding: 0.6rem 1.2rem;
        box-shadow: inset 5px 5px 12px var(--sd), inset -5px -5px 12px var(--sl);
        animation: fadeUp 0.4s ease both;
    }
    [data-testid="stExpander"] details { background: var(--bg); border:none !important; border-radius:16px;
        box-shadow: 4px 4px 9px var(--sd), -4px -4px 9px var(--sl); }
    [data-testid="stExpander"] summary { color: var(--text) !important; }

    /* ===== CHAR COUNTER ===== */
    .char-counter { font-size: 0.78rem; font-weight: 700; padding: 0.35rem 0.95rem; border-radius: 30px; display: inline-block; margin-top: 0.8rem; }
    .char-ok   { background:#dcd6ff; color:#5b21b6; }
    .char-warn { background:#ffd9d9; color:#b91c1c; }

    /* ===== NAV KARTU IKON (berbasis st.button -> tetap dalam sesi, tidak reload) ===== */
    [class*="st-key-nav_"] button {
        display: flex !important; flex-direction: column; align-items: center; justify-content: center;
        min-height: 98px; border-radius: 22px !important; border: none !important;
        background: var(--bg) !important; color: var(--soft) !important;
        font-size: 0.76rem !important; font-weight: 600 !important; line-height: 1.2;
        box-shadow: 6px 6px 13px var(--sd), -6px -6px 13px var(--sl) !important;
        transition: all 0.18s ease; animation: fadeUp 0.4s ease both;
    }
    [class*="st-key-nav_"] button:hover { transform: translateY(-3px); color: var(--accent) !important; }
    [class*="st-key-nav_"] button[kind="primary"] {
        color: var(--accent) !important;
        box-shadow: inset 5px 5px 11px var(--sd), inset -5px -5px 11px var(--sl) !important;
    }
    [class*="st-key-nav_"] button::before {
        display: flex; align-items: center; justify-content: center;
        width: 50px; height: 50px; border-radius: 50%; margin-bottom: 0.5rem; font-size: 1.45rem;
        box-shadow: 0 8px 15px rgba(40,30,90,0.32), inset 0 -4px 7px rgba(0,0,0,0.18), inset 0 4px 7px rgba(255,255,255,0.55);
        transition: transform 0.18s ease;
    }
    [class*="st-key-nav_"] button:hover::before { transform: scale(1.07); }
    .st-key-nav_ats_score    button::before { content: "▪"; background: radial-gradient(circle at 34% 28%, #6ee7b7, #059669); }
    .st-key-nav_commentary   button::before { content: "🧠"; background: radial-gradient(circle at 34% 28%, #c4b5fd, #7c3aed); }
    .st-key-nav_cv_english   button::before { content: "🌐"; background: radial-gradient(circle at 34% 28%, #93c5fd, #2563eb); }
    .st-key-nav_cv_indonesia button::before { content: "📝"; background: radial-gradient(circle at 34% 28%, #fca5a5, #ef4444); }
    .st-key-nav_cover_letter button::before { content: "✉️"; background: radial-gradient(circle at 34% 28%, #5eead4, #0d9488); }
    .st-key-nav_interview    button::before { content: "🤝"; background: radial-gradient(circle at 34% 28%, #fde68a, #f59e0b); }
    .st-key-nav_dblog        button::before { content: "💾"; background: radial-gradient(circle at 34% 28%, #f9a8d4, #db2777); }

    /* ===== ATS GAUGE ===== */
    .gauge-wrap { display:flex; align-items:center; gap:1.5rem; flex-wrap:wrap; margin:0.4rem 0 0.6rem; }
    .gauge {
        width:128px; height:128px; border-radius:50%; flex-shrink:0;
        background: conic-gradient(var(--accent) calc(var(--pct)*1%), var(--sd2) 0);
        display:flex; align-items:center; justify-content:center;
        box-shadow: 6px 6px 14px var(--sd), -6px -6px 14px var(--sl);
    }
    .gauge-inner {
        width:96px; height:96px; border-radius:50%; background:var(--bg);
        display:flex; flex-direction:column; align-items:center; justify-content:center;
        box-shadow: inset 4px 4px 9px var(--sd), inset -4px -4px 9px var(--sl);
    }
    .gauge-num { font-size:2rem; font-weight:800; color:var(--accent); line-height:1; }
    .gauge-lbl { font-size:0.62rem; color:var(--muted); margin-top:2px; }
    .kw-chip { display:inline-block; margin:3px; padding:0.3rem 0.75rem; border-radius:20px;
        background:var(--bg); box-shadow:3px 3px 7px var(--sd), -3px -3px 7px var(--sl);
        font-size:0.74rem; color:var(--soft); }

    /* ===== EMPTY STATE ===== */
    .empty-wrap { text-align:center; padding:2.4rem 1.2rem; animation: fadeUp 0.5s ease both; }
    .empty-icon { font-size:3.2rem; animation: floaty 3s ease-in-out infinite; }
    .feat-row { display:flex; gap:10px; justify-content:center; flex-wrap:wrap; margin-top:1.2rem; }
    .feat-chip { padding:0.45rem 0.9rem; border-radius:16px; background:var(--bg); font-size:0.74rem; font-weight:600;
        color:var(--soft); box-shadow:4px 4px 9px var(--sd), -4px -4px 9px var(--sl); }

    /* ===== ANIMATIONS ===== */
    @keyframes fadeUp { from {opacity:0; transform:translateY(12px);} to {opacity:1; transform:translateY(0);} }
    @keyframes floaty { 0%,100%{transform:translateY(0);} 50%{transform:translateY(-9px);} }

    .stAlert { border-radius:16px; }

    /* ===== RESPONSIVE ===== */
    @media (max-width: 768px) {
        .block-container { padding-left: 0.8rem; padding-right: 0.8rem; }
        .neo-hero { padding: 1.4rem 1.2rem; } .neo-hero h1 { font-size: 1.3rem; }
        .neo-card { padding: 1.2rem 1.1rem; }
        [class*="st-key-nav_"] button { min-height: 84px; font-size: 0.66rem !important; }
        [class*="st-key-nav_"] button::before { width: 44px; height: 44px; font-size: 1.25rem; }
    }
</style>
""", unsafe_allow_html=True)

# ══════════════════════════════════════════════════════
# STATE
# ══════════════════════════════════════════════════════
if "result" not in st.session_state:
    st.session_state.result = None
if "meta" not in st.session_state:
    st.session_state.meta = {}
if "active_view" not in st.session_state:
    st.session_state.active_view = "ats_score"

OUTPUT_CARDS = [
    ("ats_score",    "▪", "ATS Score",    "emerald"),
    ("commentary",   "🧠", "Analisis HR",  "violet"),
    ("cv_english",   "🌐", "CV English",   "blue"),
    ("cv_indonesia", "📝", "CV Indonesia", "coral"),
    ("cover_letter", "✉️", "Cover Letter", "teal"),
    ("interview",    "🤝", "Interview",    "amber"),
    ("dblog",        "💾", "Log",          "pink"),
]
VALID_VIEWS = [c[0] for c in OUTPUT_CARDS]

# ══════════════════════════════════════════════════════
# LAYOUT UTAMA
# ══════════════════════════════════════════════════════
col_in, col_out = st.columns([1, 1.35], gap="large")

with col_in:
    st.markdown("""
    <div class="neo-hero">
        <h1>📄 Buat CV Mu</h1>
        <p>AI Career Optimizer · HR · ATS · Personal Branding</p>
        <span class="pill">⚡ Powered by Gemini 2.5 Flash</span>
    </div>
    """, unsafe_allow_html=True)

    st.markdown('<div class="neo-label">📥 Input Lamaran</div>', unsafe_allow_html=True)

    uploaded_cv = st.file_uploader("1. Upload CV Anda", type=["pdf", "docx"], help="Format: PDF atau DOCX")
    uploaded_jd = st.file_uploader("2. Screenshot Lowongan (Opsional)", type=["png", "jpg", "jpeg"],
        help="Upload foto dari LinkedIn, Jobstreet, dll.")

    company_name = st.text_input("Nama Perusahaan", placeholder="Contoh: PT Tokopedia")
    position     = st.text_input("Posisi yang Dilamar", placeholder="Contoh: Senior Account Manager")
    extra_info   = st.text_area("Deskripsi Pekerjaan / Catatan Tambahan",
        placeholder="Tempel teks JD di sini jika tidak ada screenshot...", height=150)

    analyze_btn = st.button("→ Analisis & Optimalkan", type="primary", key="analyze")

# ─────────────────────────── PROSES ───────────────────────────
if analyze_btn:
    if not uploaded_cv:
        st.session_state.result = {"error": "✗ Upload CV terlebih dahulu."}
    elif not company_name.strip() or not position.strip():
        st.session_state.result = {"error": "✗ Nama perusahaan dan posisi wajib diisi."}
    elif not (manual_api_key or gemini_key):
        st.session_state.result = {"error": "✗ Gemini API Key tidak ditemukan."}
    else:
        with st.spinner("🧠 AI sedang membedah profil Anda secara mendalam..."):
            cv_text = extract_cv_text(uploaded_cv)
            if cv_text.startswith("[ERROR"):
                st.session_state.result = {"error": cv_text}
            else:
                raw_result = analyze_application(
                    cv_text=cv_text, company_name=company_name, position=position,
                    extra_info=extra_info, job_image=uploaded_jd if uploaded_jd else None,
                    api_key=manual_api_key or None
                )
                sections = parse_sections(raw_result)
                st.session_state.result = sections
                st.session_state.meta = {"company": company_name, "position": position}
                st.session_state.active_view = "ats_score"

                if sections["rejected"]:
                    save_to_supabase(company_name, position, "REJECTED", raw_result, "")
                else:
                    cl_snippet = sections["cover_letter"][:100] if sections["cover_letter"] else ""
                    save_to_supabase(company_name, position, "APPROVED", raw_result, cl_snippet)

# ─────────────────────────── OUTPUT ───────────────────────────
with col_out:
    st.markdown('<div class="neo-label">📤 Hasil Analisis</div>', unsafe_allow_html=True)
    res = st.session_state.result

    if res is None:
        st.markdown("""
        <div class="neo-card empty-wrap">
            <div class="empty-icon">🪄</div>
            <p style="font-weight:700; margin:0.8rem 0 0.2rem; font-size:1.05rem;">Siap mengoptimalkan karirmu</p>
            <p style="font-size:0.85rem; color:var(--muted);">Isi form di kiri lalu klik <b>Analisis & Optimalkan</b>.</p>
            <div class="feat-row">
                <span class="feat-chip">▪ Skor ATS</span>
                <span class="feat-chip">🧠 Analisis HR</span>
                <span class="feat-chip">🌐 CV Dwibahasa</span>
                <span class="feat-chip">✉️ Cover Letter</span>
                <span class="feat-chip">🤝 Interview Prep</span>
            </div>
        </div>
        """, unsafe_allow_html=True)

    elif "error" in res:
        st.error(res["error"])

    elif res.get("rejected"):
        st.error("! DITOLAK OLEH AI GATEKEEPER")
        with st.container(border=True):
            st.markdown(res["commentary"])

    else:
        company = st.session_state.meta.get("company", "")
        today = datetime.now().strftime("%Y%m%d")

        view = st.session_state.active_view
        if view not in VALID_VIEWS:
            view = "ats_score"

        # --- GRID KARTU IKON (st.button: navigasi tetap dalam sesi, TIDAK reload) ---
        row1 = st.columns(4, gap="small")
        row2 = st.columns(3, gap="small")
        slots = row1 + row2
        for slot, (key, icon, label, color) in zip(slots, OUTPUT_CARDS):
            with slot:
                if st.button(label, key=f"nav_{key}",
                             type="primary" if key == view else "secondary",
                             use_container_width=True):
                    st.session_state.active_view = key
                    st.rerun()

        content = res.get(view, "")

        # --- ATS SCORE ---
        if view == "ats_score":
            score, summary, matched, missing = parse_ats(res.get("ats_score", ""))
            with st.container(border=True):
                st.markdown("#### ▪ Skor Kecocokan ATS")
                st.markdown(
                    f'<div class="gauge-wrap">'
                    f'<div class="gauge" style="--pct:{score};"><div class="gauge-inner">'
                    f'<span class="gauge-num">{score}</span><span class="gauge-lbl">/ 100</span>'
                    f'</div></div>'
                    f'<div style="flex:1; min-width:200px;"><p style="margin:0; font-weight:600;">{html.escape(summary)}</p></div>'
                    f'</div>',
                    unsafe_allow_html=True
                )
                if matched:
                    st.markdown("**✓ Keyword yang sudah cocok:**")
                    st.markdown('<div>' + ''.join(f'<span class="kw-chip">{html.escape(k)}</span>' for k in matched) + '</div>', unsafe_allow_html=True)
                if missing:
                    st.markdown("**🚧 Keyword yang sebaiknya ditambahkan:**")
                    st.markdown('<div>' + ''.join(f'<span class="kw-chip">{html.escape(k)}</span>' for k in missing) + '</div>', unsafe_allow_html=True)
                if not res.get("ats_score"):
                    st.info("Skor ATS tidak terdeteksi pada output AI.")

        # --- COMMENTARY ---
        elif view == "commentary":
            with st.container(border=True):
                st.markdown("#### 🧠 Analisis Mendalam — AI HR Psikolog")
                st.markdown(content if content else res["raw"])

        # --- CV ENGLISH ---
        elif view == "cv_english":
            with st.container(border=True):
                st.markdown("#### 🌐 ATS-Optimized CV — English")
                st.markdown(content if content else "_Section tidak terdeteksi._")
            if content:
                c1, c2 = st.columns(2)
                c1.download_button("📄 Download PDF", data=make_cv_pdf(content),
                    file_name=f"CV_English_{company.replace(' ','_')}_{today}.pdf", mime="application/pdf")
                c2.download_button("📝 Download TXT", data=content.encode("utf-8"),
                    file_name=f"CV_English_{company.replace(' ','_')}_{today}.txt", mime="text/plain")

        # --- CV INDONESIA ---
        elif view == "cv_indonesia":
            with st.container(border=True):
                st.markdown("#### 📝 ATS-Optimized CV — Indonesia")
                st.markdown(content if content else "_Section tidak terdeteksi._")
            if content:
                c1, c2 = st.columns(2)
                c1.download_button("📄 Download PDF", data=make_cv_pdf(content),
                    file_name=f"CV_Indonesia_{company.replace(' ','_')}_{today}.pdf", mime="application/pdf")
                c2.download_button("📝 Download TXT", data=content.encode("utf-8"),
                    file_name=f"CV_Indonesia_{company.replace(' ','_')}_{today}.txt", mime="text/plain")

        # --- COVER LETTER ---
        elif view == "cover_letter":
            with st.container(border=True):
                st.markdown("#### ✉️ Cover Letter — Profesional & Berkarakter")
                if content:
                    st.markdown(content)
                    n = len(content)
                    cls = "char-ok" if n <= 500 else "char-warn"
                    icon = "✓" if n <= 500 else "⚠️ MELEBIHI BATAS"
                    st.markdown(f'<span class="char-counter {cls}">{icon} {n} / 500 karakter</span>', unsafe_allow_html=True)
                else:
                    st.info("_Section tidak terdeteksi._")
            if content:
                st.download_button("📄 Download Cover Letter (PDF)", data=make_pdf("Cover Letter", content),
                    file_name=f"CoverLetter_{company.replace(' ','_')}_{today}.pdf", mime="application/pdf")

        # --- INTERVIEW ---
        elif view == "interview":
            with st.container(border=True):
                st.markdown("#### 🤝 Simulasi Interview — Framework Jawaban")
                st.markdown(content if content else "_Section tidak terdeteksi._")
            if content:
                st.download_button("📄 Download Interview Prep (PDF)", data=make_pdf("Interview Preparation", content),
                    file_name=f"Interview_{company.replace(' ','_')}_{today}.pdf", mime="application/pdf")

        # --- LOG ---
        elif view == "dblog":
            with st.container(border=True):
                st.markdown("#### 💾 Database Log Entry")
                st.markdown(content if content else "_Tidak ada log._")
            with st.expander("📄 Tampilkan raw output"):
                st.text(res["raw"])
            st.download_button("📥 Download Full Output (.txt)", data=res["raw"].encode("utf-8"),
                file_name=f"FullOutput_{company.replace(' ','_')}_{datetime.now().strftime('%Y%m%d_%H%M')}.txt", mime="text/plain")

# ══════════════════════════════════════════════════════
# RIWAYAT LAMARAN
# ══════════════════════════════════════════════════════
st.markdown("---")
st.markdown('<div class="neo-label">▪ Riwayat Lamaran</div>', unsafe_allow_html=True)
if supabase:
    with st.expander("Lihat Riwayat dari Database"):
        df = fetch_history()
        if not df.empty:
            df.columns = ["Waktu", "Perusahaan", "Posisi", "Status", "Cuplikan Cover Letter"]
            st.dataframe(df, use_container_width=True)
            st.download_button("📥 Export Riwayat (.csv)", data=df.to_csv(index=False).encode("utf-8"),
                file_name=f"riwayat_{datetime.now().strftime('%Y%m%d')}.csv", mime="text/plain".replace("plain","csv"))
        else:
            st.info("Belum ada riwayat.")
else:
    st.caption("Hubungkan Supabase untuk melihat riwayat.")
